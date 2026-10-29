"""Regenerate parser fixtures in datasets/fixtures/ (synthetic Bamboo Harbor content).

    uv run python scripts/make_fixtures.py [--font /path/to/DejaVuSans.ttf]

The PDF needs a TTF font with Vietnamese glyphs (DejaVu Sans works). Outputs are committed,
so tests never depend on this script or on fonts being installed.
"""

import argparse
from datetime import UTC, datetime
from pathlib import Path

from docx import Document as DocxDocument
from fpdf import FPDF

FIXTURES = Path(__file__).resolve().parents[1] / "datasets" / "fixtures"
FIXED_DATE = datetime(2026, 1, 1, tzinfo=UTC)
DEFAULT_FONT = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")

LEAVE_POLICY_VI: list[tuple[str, list[str]]] = [
    (
        "Điều 1. Phạm vi áp dụng",
        [
            "Quy chế này áp dụng cho toàn bộ người lao động làm việc theo hợp đồng lao động "
            "tại Công ty Cổ phần Logistics Bamboo Harbor, bao gồm văn phòng Hà Nội và chi nhánh "
            "Thành phố Hồ Chí Minh.",
            "Người lao động thử việc được hưởng chế độ nghỉ phép theo tỷ lệ thời gian làm việc.",
        ],
    ),
    (
        "Điều 2. Số ngày nghỉ phép năm",
        [
            "Người lao động làm việc đủ 12 tháng được nghỉ 12 ngày làm việc mỗi năm và hưởng "
            "nguyên lương.",
            "Cứ đủ 5 năm làm việc tại công ty, số ngày nghỉ phép năm được tăng thêm 1 ngày.",
            "Ngày nghỉ phép chưa sử dụng được chuyển sang năm sau tối đa 5 ngày và phải sử dụng "
            "trước ngày 31 tháng 3.",
        ],
    ),
    (
        "Điều 3. Thủ tục xin nghỉ phép",
        [
            "Người lao động gửi đơn xin nghỉ phép trên hệ thống nhân sự ít nhất 3 ngày làm việc "
            "trước ngày nghỉ. Trường hợp nghỉ từ 5 ngày trở lên phải gửi đơn trước 10 ngày.",
            "Quản lý trực tiếp phê duyệt trong vòng 2 ngày làm việc. Phòng Nhân sự lưu hồ sơ "
            "và cập nhật số ngày phép còn lại.",
        ],
    ),
]

VPN_RUNBOOK_EN: list[tuple[str, list[str]]] = [
    (
        "Purpose",
        [
            "This runbook explains how employees connect to the Bamboo Harbor corporate network "
            "using the company VPN client. It applies to company-managed laptops only.",
        ],
    ),
    (
        "Prerequisites",
        [
            "A company-managed laptop with the VPN client version 5.2 or later.",
            "An active directory account with multi-factor authentication enrolled.",
        ],
    ),
    (
        "Connection steps",
        [
            "Open the VPN client and select the gateway closest to your office.",
            "Sign in with your directory account and approve the push notification.",
            "Confirm that the status shows Connected before opening internal tools.",
        ],
    ),
    (
        "Troubleshooting",
        [
            "If the client reports error 809, restart the network adapter and try again. If the "
            "error persists for more than 15 minutes, open a ticket with the IT service desk.",
        ],
    ),
]

VPN_GATEWAYS = [
    ("Gateway", "Region", "Address"),
    ("HN-1", "Hanoi", "vpn-hn.example.internal"),
    ("HCM-1", "Ho Chi Minh City", "vpn-hcm.example.internal"),
]

HANDBOOK_HTML = """<!doctype html>
<html lang="vi">
<head>
  <meta charset="utf-8">
  <title>Sổ tay nhân viên - Employee Handbook</title>
  <style>body { font-family: sans-serif; }</style>
  <script>window.analytics = { track: function () {} };</script>
</head>
<body>
  <header><a href="/">Bamboo Harbor Intranet</a></header>
  <nav><ul><li><a href="/hr">HR</a></li><li><a href="/it">IT</a></li></ul></nav>
  <main>
    <h1>Sổ tay nhân viên</h1>
    <p>Chào mừng bạn đến với Bamboo Harbor. Sổ tay này tóm tắt các quy định chính.</p>
    <h2>Giờ làm việc</h2>
    <p>Giờ làm việc tiêu chuẩn từ 8 giờ 30 đến 17 giờ 30, thứ Hai đến thứ Sáu,
       nghỉ trưa 1 giờ.</p>
    <h2>Working hours for engineering on-call</h2>
    <p>Engineers on the on-call rotation must respond to pages within 15 minutes,
       24 hours a day, during their on-call week.</p>
    <ul>
      <li>On-call weeks start on Monday at 09:00 local time.</li>
      <li>Compensation is one extra day off per completed on-call week.</li>
    </ul>
    <div style="display:none">Hidden tracking text that must not be indexed.</div>
    <div hidden>Another hidden element.</div>
  </main>
  <footer>© 2026 Bamboo Harbor Logistics (fictional)</footer>
</body>
</html>
"""

ONBOARDING_MD = """# Hướng dẫn hội nhập nhân viên mới

## Tuần đầu tiên

Trong tuần đầu tiên, nhân viên mới hoàn thành khóa đào tạo an toàn thông tin
và nhận laptop từ phòng IT.

## Tài khoản và công cụ

Tài khoản email và hệ thống nhân sự được cấp trong ngày làm việc đầu tiên.
Nhân viên phải đăng ký xác thực đa yếu tố trước khi truy cập VPN.
"""


def make_pdf(font: Path) -> bytes:
    pdf = FPDF()
    pdf.set_creation_date(FIXED_DATE)
    pdf.set_title("Quy chế nghỉ phép")
    pdf.add_font("DejaVu", fname=str(font))
    for page_sections in (LEAVE_POLICY_VI[:2], LEAVE_POLICY_VI[2:]):
        pdf.add_page()
        for heading, paragraphs in page_sections:
            pdf.set_font("DejaVu", size=14)
            pdf.multi_cell(0, 9, heading, new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("DejaVu", size=11)
            for paragraph in paragraphs:
                pdf.multi_cell(0, 7, paragraph, new_x="LMARGIN", new_y="NEXT")
                pdf.ln(2)
    return bytes(pdf.output())


def make_docx(path: Path) -> None:
    doc = DocxDocument()
    doc.core_properties.title = "VPN access runbook"
    doc.core_properties.created = FIXED_DATE
    doc.core_properties.modified = FIXED_DATE
    doc.add_heading("VPN access runbook", level=0)
    for heading, paragraphs in VPN_RUNBOOK_EN:
        doc.add_heading(heading, level=1)
        style = "List Number" if heading == "Connection steps" else None
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph, style=style)
        if heading == "Prerequisites":
            table = doc.add_table(rows=0, cols=3)
            for row in VPN_GATEWAYS:
                cells = table.add_row().cells
                for cell, value in zip(cells, row, strict=True):
                    cell.text = value
    doc.save(str(path))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--font", type=Path, default=DEFAULT_FONT)
    args = parser.parse_args()

    FIXTURES.mkdir(parents=True, exist_ok=True)
    (FIXTURES / "leave-policy-vi.pdf").write_bytes(make_pdf(args.font))
    make_docx(FIXTURES / "vpn-runbook-en.docx")
    (FIXTURES / "handbook-mixed.html").write_text(HANDBOOK_HTML, encoding="utf-8")
    (FIXTURES / "onboarding-vi.md").write_text(ONBOARDING_MD, encoding="utf-8")
    print(f"fixtures written to {FIXTURES}")


if __name__ == "__main__":
    main()
