"""Generate corpus v2's PDF and DOCX documents (synthetic Bamboo Harbor content, ADR-0012).

    uv run python scripts/make_corpus_v2.py [--font /path/to/DejaVuSans.ttf]

The PDF needs a TTF font with Vietnamese glyphs (DejaVu Sans or Arial Unicode work). Outputs
are committed and pinned by sha256 in the corpus manifest, so the benchmark never depends on
this script or on fonts being installed. Rerunning it is only for authoring a new corpus.
"""

# Document text is natural-language data; splitting it only hurts review.
# ruff: noqa: E501

import argparse
from datetime import UTC, datetime
from pathlib import Path

from docx import Document as DocxDocument
from fpdf import FPDF, FontFace

DOCS = Path(__file__).resolve().parents[1] / "datasets" / "corpus" / "v2" / "docs"
FIXED_DATE = datetime(2026, 10, 1, tzinfo=UTC)
DEFAULT_FONT = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")

Sections = list[tuple[str, list[str]]]

COLD_CHAIN_TITLE = "Quy trình vận chuyển hàng lạnh"
COLD_CHAIN_PAGES: list[Sections] = [
    [
        (
            "1. Phạm vi áp dụng",
            [
                "Quy trình này áp dụng cho mọi chuyến vận chuyển hàng cần kiểm soát nhiệt độ của Bamboo Harbor, từ khi nhận hàng tại kho của khách hàng đến khi giao hàng cho người nhận.",
                "Toàn bộ hàng lạnh do Bamboo Harbor vận chuyển được giữ trong khoảng 2 đến 8 độ C. Công ty không nhận vận chuyển hàng đông lạnh.",
            ],
        ),
        (
            "2. Chuẩn bị xe lạnh",
            [
                "Tài xế bật máy lạnh và làm lạnh thùng xe trước khi xếp hàng ít nhất 30 phút. Hàng chỉ được xếp lên xe khi nhiệt độ thùng xe đã nằm trong dải 2 đến 8 độ C.",
                "Thùng xe phải được vệ sinh sau mỗi chuyến và kiểm tra gioăng cửa trước ca làm việc đầu tiên trong ngày.",
            ],
        ),
        (
            "3. Giám sát nhiệt độ",
            [
                "Thiết bị ghi nhiệt độ trên xe ghi lại nhiệt độ thùng xe mỗi 15 phút và gửi dữ liệu về hệ thống điều phối.",
                "Nếu nhiệt độ nằm ngoài dải cho phép liên tục quá 20 phút, tài xế phải báo ngay cho bộ phận điều phối và làm theo hướng dẫn xử lý.",
            ],
        ),
    ],
    [
        (
            "4. Yêu cầu theo nhóm hàng",
            [
                "Bảng dưới đây quy định yêu cầu bổ sung cho từng nhóm hàng.",
            ],
        ),
        (
            "5. Giao hàng và lưu hồ sơ",
            [
                "Khi giao hàng, tài xế in hoặc gửi bản ghi nhiệt độ của chuyến cho người nhận ký xác nhận.",
                "Bản ghi nhiệt độ của mỗi chuyến được lưu trên hệ thống trong 2 năm để phục vụ kiểm tra của khách hàng và cơ quan quản lý.",
            ],
        ),
    ],
]
COLD_CHAIN_TABLE: list[tuple[str, str]] = [
    ("Nhóm hàng", "Yêu cầu bổ sung"),
    ("Vắc-xin và sinh phẩm", "Có niêm phong của khách hàng; không mở thùng trên đường."),
    ("Thuốc cần bảo quản lạnh", "Kiểm tra số lô khi nhận hàng."),
    ("Thực phẩm tươi", "Giao trong vòng 24 giờ kể từ khi nhận hàng."),
]

CODE_OF_CONDUCT_TITLE = "Bộ quy tắc ứng xử"
CODE_OF_CONDUCT: Sections = [
    (
        "Phạm vi áp dụng",
        [
            "Bộ quy tắc áp dụng cho toàn bộ người lao động, thực tập sinh và cộng tác viên của Bamboo Harbor khi làm việc với đồng nghiệp, khách hàng, nhà cung cấp và cơ quan nhà nước.",
        ],
    ),
    (
        "Quà tặng và chiêu đãi",
        [
            "Người lao động chỉ được nhận quà tặng từ khách hàng hoặc nhà cung cấp có giá trị không quá 500.000 đồng mỗi lần và không quá hai lần trong một năm từ cùng một đối tác.",
            "Không được nhận tiền mặt, thẻ quà tặng hoặc phiếu mua hàng dưới bất kỳ hình thức nào. Quà tặng vượt mức phải được từ chối lịch sự hoặc nộp lại cho phòng Hành chính.",
        ],
    ),
    (
        "Xung đột lợi ích",
        [
            "Người lao động phải khai báo xung đột lợi ích trên hệ thống nhân sự trong vòng 5 ngày làm việc kể từ khi phát sinh, ví dụ khi người thân làm việc cho nhà cung cấp hoặc khi tham gia góp vốn vào doanh nghiệp cùng ngành.",
            "Người có xung đột lợi ích không được tham gia đánh giá, lựa chọn hoặc phê duyệt hợp đồng với bên liên quan.",
        ],
    ),
    (
        "Bảo mật thông tin",
        [
            "Người lao động bảo mật thông tin của công ty và khách hàng trong thời gian làm việc và sau khi nghỉ việc. Việc phân loại và xử lý thông tin tuân theo Chính sách phân loại dữ liệu.",
        ],
    ),
    (
        "Báo cáo vi phạm",
        [
            "Người lao động báo cáo hành vi vi phạm bộ quy tắc cho quản lý trực tiếp, phòng Nhân sự hoặc qua đường dây nóng đạo đức. Đường dây nóng do một đơn vị độc lập vận hành và cho phép báo cáo ẩn danh qua điện thoại hoặc biểu mẫu trực tuyến.",
            "Công ty không xử lý kỷ luật hay trả đũa người báo cáo trung thực, kể cả khi kết quả điều tra không xác nhận vi phạm.",
        ],
    ),
]

DB_FAILOVER_TITLE = "Database Failover Runbook"
DB_FAILOVER: Sections = [
    (
        "Symptoms",
        [
            "The primary database stops accepting connections, or the database health check fails three times in a row. The API returns errors for shipment tracking and document upload.",
        ],
    ),
    (
        "Escalation",
        [
            "The primary on-call engineer declares a SEV1 incident and pages the database owner. The incident commander decides whether to fail over.",
        ],
    ),
    (
        "Steps",
        [
            "Confirm that the primary is really unavailable and not only overloaded.",
            "Check the replication lag of the standby replica. A replica may be promoted only if its replication lag is below 30 seconds.",
            "Promote the standby replica and update the database endpoint in the configuration service.",
            "Restart the API and worker services so that they connect to the new primary.",
        ],
    ),
    (
        "Recovery objectives",
        [
            "The recovery time objective (RTO) for a database failover is 15 minutes. The recovery point objective (RPO) is 1 minute of data.",
        ],
    ),
    (
        "Rollback",
        [
            "After the original primary is repaired, rebuild it as a new standby replica. Never switch back to the old primary without rebuilding it.",
        ],
    ),
]


def make_cold_chain_pdf(font: Path) -> bytes:
    pdf = FPDF()
    pdf.set_creation_date(FIXED_DATE)
    pdf.set_title(COLD_CHAIN_TITLE)
    pdf.add_font("Body", fname=str(font))
    for number, page_sections in enumerate(COLD_CHAIN_PAGES):
        pdf.add_page()
        if number == 0:
            pdf.set_font("Body", size=16)
            pdf.multi_cell(0, 10, COLD_CHAIN_TITLE, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
        for heading, paragraphs in page_sections:
            pdf.set_font("Body", size=14)
            pdf.multi_cell(0, 9, heading, new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Body", size=11)
            for paragraph in paragraphs:
                pdf.multi_cell(0, 7, paragraph, new_x="LMARGIN", new_y="NEXT")
                pdf.ln(2)
            if heading.startswith("4."):
                with pdf.table(col_widths=(60, 120), headings_style=FontFace()) as table:
                    for row in COLD_CHAIN_TABLE:
                        cells = table.row()
                        for value in row:
                            cells.cell(value)
                pdf.ln(4)
    return bytes(pdf.output())


def make_docx(path: Path, title: str, sections: Sections, *, numbered: str | None) -> None:
    doc = DocxDocument()
    doc.core_properties.title = title
    doc.core_properties.created = FIXED_DATE
    doc.core_properties.modified = FIXED_DATE
    doc.add_heading(title, level=0)
    for heading, paragraphs in sections:
        doc.add_heading(heading, level=1)
        style = "List Number" if heading == numbered else None
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph, style=style)
    doc.save(str(path))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--font", type=Path, default=DEFAULT_FONT)
    args = parser.parse_args()

    DOCS.mkdir(parents=True, exist_ok=True)
    (DOCS / "ops-cold-chain-sop.pdf").write_bytes(make_cold_chain_pdf(args.font))
    make_docx(
        DOCS / "hr-code-of-conduct.docx", CODE_OF_CONDUCT_TITLE, CODE_OF_CONDUCT, numbered=None
    )
    make_docx(
        DOCS / "eng-db-failover-runbook.docx", DB_FAILOVER_TITLE, DB_FAILOVER, numbered="Steps"
    )
    print(f"corpus v2 documents written to {DOCS}")


if __name__ == "__main__":
    main()
